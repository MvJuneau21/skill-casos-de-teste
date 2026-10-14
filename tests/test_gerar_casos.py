"""Testes do gerar_casos.py. Rodar: python -m unittest discover -s tests -v"""

import copy
import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import gerar_casos as g  # noqa: E402

BASE = {
    "funcionalidade": "Cupom no checkout",
    "criterios": [
        {"id": "CA1", "descricao": "Cupom válido aplica desconto"},
        {"id": "CA2", "descricao": "Cupom expirado não é aplicado"},
    ],
    "casos": [
        {
            "id": "CT-01", "titulo": "NEXI10 aplica 10% no Mensal", "criterios": ["CA1"],
            "tipo": "positivo", "tecnica": "caminho feliz", "prioridade": "Alta",
            "dados": {"cupom": "NEXI10"}, "passos": ["Escolher o Mensal", "Aplicar NEXI10"],
            "resultado_esperado": "Total muda de R$ 39,90 para R$ 35,91",
        },
        {
            "id": "CT-02", "titulo": "Cupom inexistente não altera o total", "criterios": ["CA1"],
            "tipo": "negativo", "prioridade": "Média",
            "dados": {"cupom": "XPTO"}, "passos": ["Escolher o Mensal", "Aplicar XPTO"],
            "resultado_esperado": "Total continua R$ 39,90",
        },
        {
            "id": "CT-03", "titulo": "VERAO15 vencido não é aplicado", "criterios": ["CA2"],
            "tipo": "negativo", "prioridade": "Média",
            "dados": {"cupom": "VERAO15"}, "passos": ["Escolher o Mensal", "Aplicar VERAO15"],
            "resultado_esperado": "Exibe \"Este cupom expirou\" e o total continua R$ 39,90",
        },
    ],
    "duvidas": [{"descricao": "Qual mensagem para cupom inexistente?", "criterio": "CA1"}],
}


def dados(**mudancas):
    d = copy.deepcopy(BASE)
    d.update(mudancas)
    return d


def tem(lista, trecho):
    return any(trecho in item for item in lista)


class Validacao(unittest.TestCase):
    def test_base_valida_sem_avisos(self):
        erros, avisos = g.validar(dados())
        self.assertEqual(erros, [])
        self.assertEqual(avisos, [])

    def test_campo_obrigatorio(self):
        erros, _ = g.validar(dados(casos=[]))
        self.assertTrue(tem(erros, "'casos'"))

    def test_criterio_inexistente_no_caso(self):
        d = dados()
        d["casos"][0]["criterios"] = ["CA9"]
        erros, _ = g.validar(d)
        self.assertTrue(tem(erros, "'CA9' não existe"))

    def test_criterio_sem_caso_e_sem_duvida_e_erro(self):
        d = dados()
        d["criterios"].append({"id": "CA3", "descricao": "Um cupom por pedido"})
        erros, _ = g.validar(d)
        self.assertTrue(tem(erros, "Critério CA3 sem nenhum caso"))

    def test_criterio_coberto_so_por_duvida_e_aviso(self):
        d = dados()
        d["criterios"].append({"id": "CA3", "descricao": "Valor mínimo R$ 1,00"})
        d["duvidas"].append({"descricao": "Que plano chega abaixo de R$ 1,00?", "criterio": ["CA3"]})
        erros, avisos = g.validar(d)
        self.assertEqual(erros, [])
        self.assertTrue(tem(avisos, "CA3 sem casos"))

    def test_duvida_com_criterio_inexistente(self):
        erros, _ = g.validar(dados(duvidas=[{"descricao": "x", "criterio": "CA7"}]))
        self.assertTrue(tem(erros, "'CA7', que não existe"))

    def test_ids_duplicados(self):
        d = dados()
        d["casos"][1]["id"] = "CT-01"
        erros, _ = g.validar(d)
        self.assertTrue(tem(erros, "'CT-01' duplicado"))

    def test_tipo_e_prioridade_invalidos(self):
        d = dados()
        d["casos"][0]["tipo"] = "feliz"
        d["casos"][0]["prioridade"] = "Urgente"
        erros, _ = g.validar(d)
        self.assertTrue(tem(erros, "tipo 'feliz'"))
        self.assertTrue(tem(erros, "prioridade 'Urgente'"))

    def test_resultado_vago(self):
        d = dados()
        d["casos"][0]["resultado_esperado"] = "O sistema funciona corretamente"
        _, avisos = g.validar(d)
        self.assertTrue(tem(avisos, "vago"))

    def test_ou_fora_das_aspas_avisa_dentro_nao(self):
        d = dados()
        d["casos"][1]["resultado_esperado"] = "Exibe erro ou ignora o cupom"
        d["casos"][2]["resultado_esperado"] = "Exibe \"E-mail ou senha incorretos\" abaixo do campo"
        _, avisos = g.validar(d)
        self.assertTrue(tem(avisos, "Caso CT-02: resultado esperado com \"ou\""))
        self.assertFalse(tem(avisos, "Caso CT-03: resultado esperado com \"ou\""))

    def test_criterio_so_positivo(self):
        d = dados()
        d["casos"][1]["tipo"] = "positivo"
        _, avisos = g.validar(d)
        self.assertTrue(tem(avisos, "CA1 só tem casos positivos"))

    def test_prioridade_inflada(self):
        d = dados()
        d["casos"].append(dict(d["casos"][2], id="CT-04", titulo="Outro", dados={"cupom": "X"}))
        for k in d["casos"]:
            k["prioridade"] = "Alta"
        _, avisos = g.validar(d)
        self.assertTrue(tem(avisos, "4 de 4 casos estão com prioridade Alta"))

    def test_casos_duplicados(self):
        d = dados()
        d["casos"].append(dict(d["casos"][0], id="CT-04", titulo="Outro título"))
        _, avisos = g.validar(d)
        self.assertTrue(tem(avisos, "CT-01 e CT-04 têm os mesmos passos e dados"))

    def test_caso_longo(self):
        d = dados()
        d["casos"][0]["passos"] = [f"Passo {i}" for i in range(12)]
        _, avisos = g.validar(d)
        self.assertTrue(tem(avisos, "12 passos"))


    def test_normalmente_e_vago(self):
        d = dados()
        d["casos"][0]["resultado_esperado"] = "O título reproduz normalmente"
        _, avisos = g.validar(d)
        self.assertTrue(tem(avisos, "vago"))

    def test_bloqueia_e_suposicao_com_caso_inexistente(self):
        d = dados()
        d["duvidas"][0]["bloqueia"] = ["CT-99"]
        d["suposicoes"] = [{"descricao": "x", "casos": ["CT-98"]}]
        erros, _ = g.validar(d)
        self.assertTrue(tem(erros, "bloqueia o caso 'CT-99'"))
        self.assertTrue(tem(erros, "caso 'CT-98'"))

    def test_muitas_duvidas_sem_bloqueia(self):
        d = dados(duvidas=[{"descricao": f"Dúvida {i}"} for i in range(7)])
        _, avisos = g.validar(d)
        self.assertTrue(tem(avisos, "7 dúvidas e nenhuma com 'bloqueia'"))
        d["duvidas"][0]["bloqueia"] = ["CT-02"]
        _, avisos = g.validar(d)
        self.assertFalse(tem(avisos, "nenhuma com 'bloqueia'"))


    def test_duvidas_demais(self):
        d = dados(duvidas=[{"descricao": f"Dúvida {i}", "bloqueia": ["CT-02"]} for i in range(9)])
        _, avisos = g.validar(d)
        self.assertTrue(tem(avisos, "9 dúvidas (limite 8)"))
        d["duvidas"].pop()
        _, avisos = g.validar(d)
        self.assertFalse(tem(avisos, "limite 8"))


    def test_ou_em_lista_negada_nao_avisa(self):  # v1.4
        d = dados()
        d["casos"][1]["resultado_esperado"] = "A resposta não pede senha, código ou cartão"
        _, avisos = g.validar(d)
        self.assertFalse(tem(avisos, "com \"ou\""))

    def test_ou_em_caso_bloqueado_nao_avisa(self):  # v1.4
        d = dados()
        d["casos"][1]["resultado_esperado"] = "Exibe erro ou ignora o cupom, conforme a resposta do PO"
        d["duvidas"][0]["bloqueia"] = ["CT-02"]
        _, avisos = g.validar(d)
        self.assertFalse(tem(avisos, "com \"ou\""))

    def test_resultado_curto_sugere_efeito(self):  # v1.4
        d = dados()
        d["casos"][1]["resultado_esperado"] = "Status 400"
        _, avisos = g.validar(d)
        self.assertTrue(tem(avisos, "nada é gravado"))


class Geracao(unittest.TestCase):
    def test_gera_md_e_xlsx(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            arq = Path(tmp) / "casos.json"
            arq.write_text(json.dumps(BASE, ensure_ascii=False), encoding="utf-8")
            self.assertEqual(g.main([str(arq), "--saida", tmp]), 0)

            md = (Path(tmp) / "cupom-no-checkout_casos-de-teste.md").read_text(encoding="utf-8")
            self.assertIn("### CT-03 · VERAO15 vencido não é aplicado", md)
            self.assertIn("**D1.** Qual mensagem para cupom inexistente?", md)
            self.assertIn("| CA1 | Cupom válido aplica desconto | CT-01, CT-02 | D1 |", md)

            wb = load_workbook(Path(tmp) / "cupom-no-checkout_casos-de-teste.xlsx")
            self.assertEqual(wb.sheetnames, ["Resumo", "Casos", "Rastreabilidade", "Dúvidas"])
            self.assertEqual(wb["Casos"]["L2"].value, "Não executado")
            self.assertEqual(wb["Rastreabilidade"]["C2"].value, "X")
            self.assertEqual(wb["Rastreabilidade"].cell(row=3, column=6).value, 1)

    def test_caso_bloqueado_e_dependencias(self):
        from openpyxl import load_workbook

        d = dados(suposicoes=[{"descricao": "Validade inclui o último dia", "casos": ["CT-03"]}])
        d["duvidas"][0]["bloqueia"] = ["CT-02"]
        with tempfile.TemporaryDirectory() as tmp:
            arq = Path(tmp) / "casos.json"
            arq.write_text(json.dumps(d, ensure_ascii=False), encoding="utf-8")
            self.assertEqual(g.main([str(arq), "--saida", tmp]), 0)
            md = (Path(tmp) / "cupom-no-checkout_casos-de-teste.md").read_text(encoding="utf-8")
            self.assertIn("**Dúvidas que bloqueiam a execução**", md)
            self.assertIn("> Depende de D1 (bloqueia)", md)
            ws = load_workbook(Path(tmp) / "cupom-no-checkout_casos-de-teste.xlsx")["Casos"]
            self.assertEqual(ws["K3"].value, "D1 (bloqueia)")
            self.assertEqual(ws["L3"].value, "Bloqueado")
            self.assertEqual(ws["K4"].value, "Suposição: Validade inclui o último dia")
            self.assertEqual(ws["L4"].value, "Não executado")

    def test_registrar_execucao_na_planilha(self):  # v1.4
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            arq = Path(tmp) / "casos.json"
            arq.write_text(json.dumps(BASE, ensure_ascii=False), encoding="utf-8")
            self.assertEqual(g.main([str(arq), "--saida", tmp]), 0)
            res = Path(tmp) / "resultados.json"
            res.write_text(json.dumps({
                "executado_por": "Camila", "data": "14/10/2026",
                "casos": {"CT-01": {"status": "Passou", "obtido": "Total R$ 35,91"},
                          "CT-03": {"status": "Falhou", "obtido": "Total mudou", "bug": "BUG-7"}},
                "duvidas": {"D1": "Mensagem: Cupom inválido"}}, ensure_ascii=False), encoding="utf-8")
            self.assertEqual(g.main([str(arq), "--saida", tmp, "--registrar", str(res)]), 0)

            wb = load_workbook(Path(tmp) / "cupom-no-checkout_casos-de-teste.xlsx")
            ws = wb["Casos"]
            col = {c.value: c.column for c in ws[1]}
            self.assertEqual(ws.cell(2, col["Status"]).value, "Passou")
            self.assertEqual(ws.cell(2, col["Resultado obtido"]).value, "Total R$ 35,91")
            self.assertEqual(ws.cell(4, col["Status"]).value, "Falhou")
            self.assertEqual(ws.cell(4, col["Bug / ticket"]).value, "BUG-7")
            self.assertEqual(ws.cell(4, col["Executado por"]).value, "Camila")
            self.assertEqual(ws.cell(3, col["Status"]).value, "Não executado")  # CT-02 sem resultado
            self.assertEqual(wb["Dúvidas"]["E2"].value, "Mensagem: Cupom inválido")

    def test_registrar_recusa_caso_ou_status_invalido(self):  # v1.4
        with tempfile.TemporaryDirectory() as tmp:
            arq = Path(tmp) / "casos.json"
            arq.write_text(json.dumps(BASE, ensure_ascii=False), encoding="utf-8")
            g.main([str(arq), "--saida", tmp])
            res = Path(tmp) / "resultados.json"
            res.write_text(json.dumps({"casos": {"CT-99": {"status": "Passou"}, "CT-01": {"status": "OK"}}}),
                           encoding="utf-8")
            self.assertEqual(g.main([str(arq), "--saida", tmp, "--registrar", str(res)]), 1)

    def test_registrar_sem_planilha(self):  # v1.4
        with tempfile.TemporaryDirectory() as tmp:
            arq = Path(tmp) / "casos.json"
            arq.write_text(json.dumps(BASE, ensure_ascii=False), encoding="utf-8")
            res = Path(tmp) / "resultados.json"
            res.write_text("{}", encoding="utf-8")
            self.assertEqual(g.main([str(arq), "--saida", tmp, "--registrar", str(res)]), 2)

    def test_erro_nao_gera_nada(self):
        with tempfile.TemporaryDirectory() as tmp:
            arq = Path(tmp) / "casos.json"
            arq.write_text(json.dumps(dados(criterios=[])), encoding="utf-8")
            self.assertEqual(g.main([str(arq), "--saida", tmp]), 1)
            self.assertEqual(list(Path(tmp).glob("*.md")), [])

    def test_json_invalido(self):
        with tempfile.TemporaryDirectory() as tmp:
            arq = Path(tmp) / "casos.json"
            arq.write_text("{ quebrado", encoding="utf-8")
            self.assertEqual(g.main([str(arq)]), 2)


if __name__ == "__main__":
    unittest.main()
