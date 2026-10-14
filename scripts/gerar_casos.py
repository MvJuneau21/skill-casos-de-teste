#!/usr/bin/env python3
"""
Valida o JSON de casos de teste e gera os arquivos finais no padrão Nexi:
  - <funcionalidade>_casos-de-teste.md   -> documento para revisar com PO/dev
  - <funcionalidade>_casos-de-teste.xlsx -> planilha de execução (status, bug, evidência)

Validação (o script só gera os arquivos se não houver ERROS):
  ERRO  - campo obrigatório ausente, ID duplicado, caso ligado a critério que
          não existe, tipo/prioridade inválidos, critério sem nenhum caso e
          sem dúvida registrada (buraco de cobertura), dúvida ou suposição
          apontando para caso que não existe.
  AVISO - resultado esperado vago ou com "ou", critério só com casos positivos,
          critério coberto só por dúvida, prioridade inflada, caso longo demais,
          casos duplicados, técnica fora da lista, mais de 8 dúvidas, muitas
          dúvidas sem dizer quais bloqueiam a execução.

Uso:
  python gerar_casos.py casos.json                    # valida e gera em ./
  python gerar_casos.py casos.json --saida ./qa       # escolhe a pasta
  python gerar_casos.py casos.json --validar          # só valida, não gera
  python gerar_casos.py casos.json --saida ./qa --registrar resultados.json
                                                      # grava a execução na planilha já gerada

Formato do resultados.json (todos os campos de cada caso são opcionais, exceto status):
  {"executado_por": "Ana", "data": "14/10/2026",
   "casos": {"CT-01": {"status": "Passou", "obtido": "...", "evidencia": "...", "bug": "..."}},
   "duvidas": {"D1": "resposta do PO ou da execução"}}

Códigos de saída: 0 = ok, 1 = erros de validação, 2 = erro de uso/arquivo.
"""

import argparse
import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

TIPOS = ["positivo", "negativo", "borda"]
PRIORIDADES = ["Alta", "Média", "Baixa"]
FONTES = ["explícito", "derivado"]
TECNICAS = [
    "caminho feliz",
    "partição de equivalência",
    "valor limite",
    "tabela de decisão",
    "transição de estado",
    "suposição de erro",
]
STATUS_EXEC = ["Não executado", "Passou", "Falhou", "Bloqueado"]
MAX_PASSOS = 10
MAX_DUVIDAS_SEM_BLOQUEIO = 6
MAX_DUVIDAS = 8

# Frases que não dizem o que observar. Texto entre aspas (mensagens da tela) é ignorado.
VAGO = re.compile(
    r"\b(funciona(r)? (corretamente|normalmente|bem)|como esperado|comportamento "
    r"(esperado|correto|adequado)|sistema responde (corretamente|bem)|tudo certo|"
    r"sem problemas?|de forma (correta|adequada)|normalmente|corretamente|adequadamente)\b",
    re.IGNORECASE,
)
ASPAS = re.compile(r"\"[^\"]*\"|'[^']*'|“[^”]*”")


def normalize(text):
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", text).strip().lower()


def slugify(text, limit=50):
    return re.sub(r"[^a-z0-9]+", "-", normalize(text)).strip("-")[:limit].strip("-") or "funcionalidade"


def shorten(text, limit=45):
    """Corta em palavra inteira, para o aviso não terminar no meio de uma palavra."""
    text = str(text)
    return text if len(text) <= limit else text[:limit].rsplit(" ", 1)[0] + "…"


def ou_indefinido(texto):
    """True se algum trecho tem "ou" sem uma negação antes dele ("não pede A ou B" é uma lista)."""
    for trecho in re.split(r"[;.]", normalize(texto)):
        pos = re.search(r"\bou\b", trecho)
        if pos and not re.search(r"\b(nao|nenhum|nenhuma|nem|sem)\b", trecho[:pos.start()]):
            return True
    return False


def as_list(value):
    if value is None:
        return []
    return value if isinstance(value, list) else [value]


# --------------------------------------------------------------------------- #
# Validação
# --------------------------------------------------------------------------- #
def validar(dados):
    erros, avisos = [], []

    if not isinstance(dados, dict):
        return ["O JSON precisa ser um objeto { ... }."], avisos

    for campo in ("funcionalidade", "criterios", "casos"):
        if not dados.get(campo):
            erros.append(f"Campo obrigatório ausente ou vazio: '{campo}'.")
    if erros:
        return erros, avisos

    # Critérios
    crit_ids = []
    for i, c in enumerate(as_list(dados["criterios"]), 1):
        if not isinstance(c, dict) or not c.get("id") or not c.get("descricao"):
            erros.append(f"Critério nº {i}: precisa de 'id' e 'descricao'.")
            continue
        if c["id"] in crit_ids:
            erros.append(f"Critério '{c['id']}' duplicado.")
        crit_ids.append(c["id"])
        fonte = c.get("fonte", "explícito")
        if fonte not in FONTES:
            erros.append(f"Critério {c['id']}: fonte '{fonte}' inválida (use {' ou '.join(FONTES)}).")

    # Dúvidas (validadas antes da cobertura, porque uma dúvida pode "cobrir" um critério)
    crit_com_duvida, bloqueados = set(), set()
    for i, d in enumerate(as_list(dados.get("duvidas")), 1):
        if not isinstance(d, dict) or not d.get("descricao"):
            erros.append(f"Dúvida nº {i}: precisa de 'descricao'.")
            continue
        for cid in as_list(d.get("criterio")):
            if cid not in crit_ids:
                erros.append(f"Dúvida nº {i} aponta para o critério '{cid}', que não existe.")
            crit_com_duvida.add(cid)
        bloqueados.update(as_list(d.get("bloqueia")))

    # Casos
    caso_ids, titulos, assinaturas = [], {}, {}
    por_criterio = {cid: [] for cid in crit_ids}
    altas = 0
    casos = as_list(dados["casos"])
    for i, k in enumerate(casos, 1):
        if not isinstance(k, dict):
            erros.append(f"Caso nº {i}: precisa ser um objeto.")
            continue
        cid = k.get("id") or f"nº {i}"
        for campo in ("id", "titulo", "criterios", "tipo", "prioridade", "passos", "resultado_esperado"):
            if not k.get(campo):
                erros.append(f"Caso {cid}: campo obrigatório '{campo}' ausente ou vazio.")
        if k.get("id"):
            if k["id"] in caso_ids:
                erros.append(f"Caso '{k['id']}' duplicado.")
            caso_ids.append(k["id"])

        for crit in as_list(k.get("criterios")):
            if crit not in por_criterio:
                erros.append(f"Caso {cid}: critério '{crit}' não existe em 'criterios'.")
            else:
                por_criterio[crit].append(k.get("tipo"))

        if k.get("tipo") and k["tipo"] not in TIPOS:
            erros.append(f"Caso {cid}: tipo '{k['tipo']}' inválido (use {', '.join(TIPOS)}).")
        if k.get("prioridade") and k["prioridade"] not in PRIORIDADES:
            erros.append(f"Caso {cid}: prioridade '{k['prioridade']}' inválida (use {', '.join(PRIORIDADES)}).")
        if k.get("prioridade") == "Alta":
            altas += 1
        if k.get("tecnica") and normalize(k["tecnica"]) not in {normalize(t) for t in TECNICAS}:
            avisos.append(f"Caso {cid}: técnica '{k['tecnica']}' fora da lista ({', '.join(TECNICAS)}).")

        passos = as_list(k.get("passos"))
        if len(passos) > MAX_PASSOS:
            avisos.append(f"Caso {cid}: {len(passos)} passos. Caso grande costuma testar mais de uma coisa; considere dividir.")

        resultado = str(k.get("resultado_esperado") or "")
        sem_aspas = ASPAS.sub("", resultado)
        if resultado and VAGO.search(sem_aspas):
            avisos.append(f"Caso {cid}: resultado esperado vago (\"{shorten(resultado)}\"). Diga o que se vê na tela/dado.")
        elif resultado and len(sem_aspas.split()) < 3 and not ASPAS.search(resultado):
            avisos.append(
                f"Caso {cid}: resultado esperado curto demais (\"{resultado}\"). Em API ou CLI, diga também "
                "o efeito (ex.: \"Status 400 e nada é gravado\")."
            )
        # Caso bloqueado depende da resposta do PO, então o "ou" é esperado; "não pede A ou B" é lista, não dúvida.
        if k.get("id") not in bloqueados and ou_indefinido(sem_aspas):
            avisos.append(
                f"Caso {cid}: resultado esperado com \"ou\" (\"{shorten(resultado)}\"). "
                "Se o comportamento não está definido, registre em 'duvidas'."
            )

        if k.get("titulo"):
            t = normalize(k["titulo"])
            if t in titulos:
                avisos.append(f"Casos {titulos[t]} e {cid} têm o mesmo título.")
            titulos.setdefault(t, cid)
        assinatura = normalize(json.dumps([passos, k.get("dados")], ensure_ascii=False, sort_keys=True))
        if passos and assinatura in assinaturas:
            avisos.append(f"Casos {assinaturas[assinatura]} e {cid} têm os mesmos passos e dados.")
        assinaturas.setdefault(assinatura, cid)

    # Dúvidas e suposições que apontam para casos
    for i, d in enumerate(as_list(dados.get("duvidas")), 1):
        for ref in as_list(d.get("bloqueia")) if isinstance(d, dict) else []:
            if ref not in caso_ids:
                erros.append(f"Dúvida nº {i} bloqueia o caso '{ref}', que não existe.")
    for i, sup in enumerate(as_list(dados.get("suposicoes")), 1):
        for ref in as_list(sup.get("casos")) if isinstance(sup, dict) else []:
            if ref not in caso_ids:
                erros.append(f"Suposição nº {i} aponta para o caso '{ref}', que não existe.")
    duvidas = [d for d in as_list(dados.get("duvidas")) if isinstance(d, dict)]
    if len(duvidas) > MAX_DUVIDAS_SEM_BLOQUEIO and not any(d.get("bloqueia") for d in duvidas):
        avisos.append(
            f"{len(duvidas)} dúvidas e nenhuma com 'bloqueia'. Marque quais casos não rodam sem a resposta, "
            "para o PO saber o que responder primeiro."
        )
    if len(duvidas) > MAX_DUVIDAS:
        avisos.append(
            f"{len(duvidas)} dúvidas (limite {MAX_DUVIDAS}). Junte as parecidas numa só (ex.: todas as mensagens "
            "que o requisito não define); lista longa não é lida pelo PO."
        )

    # Cobertura
    for crit, tipos in por_criterio.items():
        if not tipos:
            if crit in crit_com_duvida:
                avisos.append(f"Critério {crit} sem casos: coberto só por dúvida. Ele fica sem teste até a resposta do PO.")
            else:
                erros.append(
                    f"Critério {crit} sem nenhum caso. Crie um caso ou, se falta informação para testar, "
                    f"registre uma dúvida com \"criterio\": \"{crit}\"."
                )
        elif all(t == "positivo" for t in tipos):
            avisos.append(f"Critério {crit} só tem casos positivos. Falta o caso negativo ou de borda?")

    if len(casos) >= 4 and altas > len(casos) / 2:
        avisos.append(
            f"{altas} de {len(casos)} casos estão com prioridade Alta. Prioridade só ajuda a ordenar "
            "a execução quando diferencia os casos; reveja."
        )

    return erros, avisos


# --------------------------------------------------------------------------- #
# Markdown
# --------------------------------------------------------------------------- #
def fmt_dados(valor):
    if isinstance(valor, dict):
        return [f"`{k}`: {v}" for k, v in valor.items()]
    return [str(v) for v in as_list(valor)]


def casos_por_criterio(dados):
    mapa = {c["id"]: [] for c in dados["criterios"]}
    for k in dados["casos"]:
        for crit in as_list(k["criterios"]):
            mapa[crit].append(k["id"])
    return mapa


def duvidas_por_criterio(dados):
    mapa = {}
    for n, d in enumerate(as_list(dados.get("duvidas")), 1):
        for crit in as_list(d.get("criterio")):
            mapa.setdefault(crit, []).append(f"D{n}")
    return mapa


def dependencias(dados):
    """Caso -> lista do que ele depende (dúvida que bloqueia, suposição)."""
    dep = {}
    for n, d in enumerate(as_list(dados.get("duvidas")), 1):
        for ref in as_list(d.get("bloqueia")):
            dep.setdefault(ref, []).append(f"D{n} (bloqueia)")
    for s in as_list(dados.get("suposicoes")):
        if isinstance(s, dict):
            for ref in as_list(s.get("casos")):
                dep.setdefault(ref, []).append(f"Suposição: {s.get('descricao', '')}")
    return dep


def contar(casos, campo, valores):
    return {v: sum(1 for k in casos if k.get(campo) == v) for v in valores}


def gerar_md(dados):
    casos = dados["casos"]
    tipos = contar(casos, "tipo", TIPOS)
    prios = contar(casos, "prioridade", PRIORIDADES)
    cobertura = casos_por_criterio(dados)
    duv_crit = duvidas_por_criterio(dados)
    duvidas = as_list(dados.get("duvidas"))
    suposicoes = as_list(dados.get("suposicoes"))

    L = [f"# Casos de teste: {dados['funcionalidade']}", ""]
    meta = []
    if dados.get("origem"):
        meta.append(f"**Origem:** {dados['origem']}")
    meta.append(f"**Gerado em:** {date.today().strftime('%d/%m/%Y')}")
    L.append(" · ".join(meta))
    L.append("")
    L.append(
        f"**{len(casos)} casos:** {tipos['positivo']} positivos, {tipos['negativo']} negativos, "
        f"{tipos['borda']} de borda · **Prioridade:** {prios['Alta']} Alta, {prios['Média']} Média, "
        f"{prios['Baixa']} Baixa"
    )
    L.append("")

    if duvidas or suposicoes:
        L += ["## Pontos a confirmar com o PO", ""]
        if duvidas:
            numeradas = list(enumerate(duvidas, 1))
            grupos = [
                ("**Dúvidas que bloqueiam a execução** (responder primeiro):", [x for x in numeradas if x[1].get("bloqueia")]),
                ("**Outras dúvidas** (o requisito não diz; nenhum caso inventa a resposta):", [x for x in numeradas if not x[1].get("bloqueia")]),
            ]
            for titulo, itens in grupos:
                if not itens:
                    continue
                L += [titulo, ""]
                for n, d in itens:
                    extra = []
                    if d.get("criterio"):
                        extra.append(", ".join(as_list(d["criterio"])))
                    if d.get("bloqueia"):
                        extra.append("bloqueia " + ", ".join(as_list(d["bloqueia"])))
                    if d.get("impacto"):
                        extra.append(d["impacto"])
                    sufixo = f" _({' · '.join(extra)})_" if extra else ""
                    L.append(f"- **D{n}.** {d['descricao']}{sufixo}")
                L.append("")
        if suposicoes:
            L += ["**Suposições** (usadas nos casos; se estiverem erradas, os casos indicados mudam):", ""]
            for s in suposicoes:
                if isinstance(s, dict):
                    casos_s = ", ".join(as_list(s.get("casos")))
                    L.append(f"- {s.get('descricao', '')}" + (f" _({casos_s})_" if casos_s else ""))
                else:
                    L.append(f"- {s}")
            L.append("")

    L += ["## Critérios e cobertura", "", "| Critério | Descrição | Casos | Dúvidas |", "|---|---|---|---|"]
    for c in dados["criterios"]:
        nome = c["id"] + (" _(derivado)_" if c.get("fonte") == "derivado" else "")
        L.append(
            f"| {nome} | {c['descricao']} | {', '.join(cobertura[c['id']]) or '**sem caso**'} | "
            f"{', '.join(duv_crit.get(c['id'], [])) or '—'} |"
        )
    L.append("")

    L += ["## Casos", ""]
    dep = dependencias(dados)
    for k in casos:
        L.append(f"### {k['id']} · {k['titulo']}")
        L.append("")
        tec = k.get("tecnica") or "—"
        L.append(
            f"**Critérios:** {', '.join(as_list(k['criterios']))} · **Tipo:** {k['tipo']} · "
            f"**Técnica:** {tec} · **Prioridade:** {k['prioridade']}"
        )
        L.append("")
        for item in dep.get(k["id"], []):
            L.append(f"> Depende de {item}")
            L.append("")
        if k.get("pre_condicoes"):
            L.append("**Pré-condições:**")
            L += [f"- {p}" for p in as_list(k["pre_condicoes"])]
            L.append("")
        if k.get("dados"):
            L.append("**Dados:**")
            L += [f"- {p}" for p in fmt_dados(k["dados"])]
            L.append("")
        L.append("**Passos:**")
        L += [f"{n}. {p}" for n, p in enumerate(as_list(k["passos"]), 1)]
        L.append("")
        L.append(f"**Resultado esperado:** {k['resultado_esperado']}")
        L.append("")

    if dados.get("fora_do_escopo"):
        L += ["## Fora do escopo", ""]
        L += [f"- {f}" for f in as_list(dados["fora_do_escopo"])]
        L.append("")

    return "\n".join(L)


# --------------------------------------------------------------------------- #
# Planilha
# --------------------------------------------------------------------------- #
def gerar_xlsx(dados, caminho):
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    cab_fill = PatternFill("solid", fgColor="1F3864")
    cab_font = Font(bold=True, color="FFFFFF")
    quebra = Alignment(wrap_text=True, vertical="top")

    def cabecalho(ws, colunas):
        ws.append([c for c, _ in colunas])
        for i, (_, largura) in enumerate(colunas, 1):
            cel = ws.cell(row=1, column=i)
            cel.fill, cel.font = cab_fill, cab_font
            cel.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = largura
        ws.freeze_panes = "A2"

    wb = Workbook()

    # Casos
    ws = wb.active
    ws.title = "Casos"
    colunas = [
        ("ID", 9), ("Título", 34), ("Critérios", 11), ("Tipo", 10), ("Técnica", 18),
        ("Prioridade", 10), ("Pré-condições", 32), ("Dados", 28), ("Passos", 46),
        ("Resultado esperado", 40), ("Depende de", 28), ("Status", 14), ("Resultado obtido", 30),
        ("Evidência (print/vídeo/log)", 24), ("Bug / ticket", 14), ("Executado por", 16), ("Data execução", 13),
    ]
    dep = dependencias(dados)
    cabecalho(ws, colunas)
    for k in dados["casos"]:
        ws.append([
            k["id"], k["titulo"], ", ".join(as_list(k["criterios"])), k["tipo"],
            k.get("tecnica", ""), k["prioridade"],
            "\n".join(f"• {p}" for p in as_list(k.get("pre_condicoes"))),
            "\n".join(fmt_dados(k.get("dados"))).replace("`", ""),
            "\n".join(f"{n}. {p}" for n, p in enumerate(as_list(k["passos"]), 1)),
            k["resultado_esperado"], "\n".join(dep.get(k["id"], [])),
            "Bloqueado" if any(x.endswith("(bloqueia)") for x in dep.get(k["id"], [])) else "Não executado",
            "", "", "", "", "",
        ])
    ultima = ws.max_row
    for row in ws.iter_rows(min_row=2, max_row=ultima):
        for cel in row:
            cel.alignment = quebra
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_EXEC) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"L2:L{max(ultima, 2)}")
    cores = {"Passou": "C6EFCE", "Falhou": "FFC7CE", "Bloqueado": "FFEB9C"}
    for status, cor in cores.items():
        ws.conditional_formatting.add(
            f"L2:L{max(ultima, 2)}",
            CellIsRule(operator="equal", formula=[f'"{status}"'], fill=PatternFill("solid", fgColor=cor)),
        )
    ws.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{ultima}"

    # Rastreabilidade
    wr = wb.create_sheet("Rastreabilidade")
    ids = [k["id"] for k in dados["casos"]]
    cabecalho(wr, [("Critério", 10), ("Descrição", 50)] + [(i, 7) for i in ids] + [("Nº casos", 9)])
    for c in dados["criterios"]:
        marcados = ["X" if c["id"] in as_list(k["criterios"]) else "" for k in dados["casos"]]
        wr.append([c["id"], c["descricao"]] + marcados + [marcados.count("X")])
    for row in wr.iter_rows(min_row=2, max_row=wr.max_row):
        row[1].alignment = quebra
        for cel in row[2:]:
            cel.alignment = Alignment(horizontal="center", vertical="top")
    col_n = get_column_letter(len(ids) + 3)
    wr.conditional_formatting.add(
        f"{col_n}2:{col_n}{max(wr.max_row, 2)}",
        CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", fgColor="FFC7CE")),
    )

    # Dúvidas
    wd = wb.create_sheet("Dúvidas")
    cabecalho(wd, [("Nº", 6), ("Dúvida", 60), ("Critério", 12), ("Impacto", 30), ("Resposta do PO", 40), ("Respondido por", 16)])
    for n, d in enumerate(as_list(dados.get("duvidas")), 1):
        wd.append([f"D{n}", d["descricao"], ", ".join(as_list(d.get("criterio"))), d.get("impacto", ""), "", ""])
    for row in wd.iter_rows(min_row=2, max_row=wd.max_row):
        for cel in row:
            cel.alignment = quebra

    # Resumo (fórmulas: atualiza sozinho conforme a execução)
    wsr = wb.create_sheet("Resumo", 0)
    wsr.column_dimensions["A"].width = 26
    wsr.column_dimensions["B"].width = 12
    wsr.append([dados["funcionalidade"]])
    wsr["A1"].font = Font(bold=True, size=14)
    wsr.append([])
    wsr.append(["Status", "Casos"])
    for status in STATUS_EXEC:
        wsr.append([status, f'=COUNTIF(Casos!L2:L{max(ultima, 2)},"{status}")'])
    wsr.append(["Total", f"=COUNTA(Casos!A2:A{max(ultima, 2)})"])
    wsr.append(["% executado", f"=IF(B8=0,0,(B8-B4)/B8)"])
    wsr["B9"].number_format = "0%"
    wsr.append([])
    wsr.append(["Falhas em prioridade Alta", f'=COUNTIFS(Casos!L2:L{max(ultima, 2)},"Falhou",Casos!F2:F{max(ultima, 2)},"Alta")'])
    wsr.append(["Dúvidas abertas", "=COUNTA('Dúvidas'!A2:A500)-COUNTA('Dúvidas'!E2:E500)"])
    for cel in ("A3", "B3"):
        wsr[cel].fill, wsr[cel].font = cab_fill, cab_font

    wb.save(caminho)


# --------------------------------------------------------------------------- #
# Registro da execução
# --------------------------------------------------------------------------- #
COLUNAS_RESULTADO = {"status": "Status", "obtido": "Resultado obtido",
                     "evidencia": "Evidência (print/vídeo/log)", "bug": "Bug / ticket"}


def registrar(dados, xlsx, resultados):
    """Grava status, resultado obtido, evidência e bug na planilha. Devolve (erros, contagem por status)."""
    from openpyxl import load_workbook

    erros = []
    ids = {k["id"] for k in dados["casos"]}
    casos = resultados.get("casos") or {}
    for cid, r in casos.items():
        if cid not in ids:
            erros.append(f"Resultado para o caso '{cid}', que não existe no casos.json.")
        elif not isinstance(r, dict) or r.get("status") not in STATUS_EXEC:
            erros.append(f"Caso {cid}: status inválido (use {', '.join(STATUS_EXEC)}).")
    n_duvidas = len(as_list(dados.get("duvidas")))
    for did in (resultados.get("duvidas") or {}):
        if not re.fullmatch(r"D\d+", str(did)) or not 1 <= int(did[1:]) <= n_duvidas:
            erros.append(f"Resposta para a dúvida '{did}', que não existe.")
    if erros:
        return erros, {}

    wb = load_workbook(xlsx)
    ws = wb["Casos"]
    col = {c.value: c.column for c in ws[1]}
    for row in ws.iter_rows(min_row=2):
        r = casos.get(row[0].value)
        if not r:
            continue
        for chave, titulo in COLUNAS_RESULTADO.items():
            if r.get(chave):
                ws.cell(row=row[0].row, column=col[titulo], value=r[chave])
        ws.cell(row=row[0].row, column=col["Executado por"], value=resultados.get("executado_por", ""))
        ws.cell(row=row[0].row, column=col["Data execução"], value=resultados.get("data", date.today().strftime("%d/%m/%Y")))
    wd = wb["Dúvidas"]
    for row in wd.iter_rows(min_row=2):
        resposta = (resultados.get("duvidas") or {}).get(row[0].value)
        if resposta:
            row[4].value = resposta
            row[5].value = resultados.get("executado_por", "")
    wb.save(xlsx)
    contagem = {s: sum(1 for r in casos.values() if r["status"] == s) for s in STATUS_EXEC}
    return [], contagem


# --------------------------------------------------------------------------- #
def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except (AttributeError, ValueError):
            pass

    ap = argparse.ArgumentParser(description="Valida o casos.json e gera o .md e o .xlsx.")
    ap.add_argument("json", help="caminho do casos.json")
    ap.add_argument("--saida", default=".", help="pasta onde gerar os arquivos (padrão: ./)")
    ap.add_argument("--validar", action="store_true", help="só valida, não gera arquivos")
    ap.add_argument("--registrar", metavar="RESULTADOS", help="grava a execução (resultados.json) na planilha já gerada")
    args = ap.parse_args(argv)

    caminho = Path(args.json)
    try:
        dados = json.loads(caminho.read_text(encoding="utf-8"))
    except FileNotFoundError:
        print(f"ERRO: arquivo não encontrado: {caminho}", file=sys.stderr)
        return 2
    except json.JSONDecodeError as e:
        print(f"ERRO: JSON inválido em {caminho}: linha {e.lineno}, coluna {e.colno}: {e.msg}", file=sys.stderr)
        return 2

    erros, avisos = validar(dados)
    for e in erros:
        print(f"ERRO: {e}")
    for a in avisos:
        print(f"AVISO: {a}")
    if erros:
        print(f"\n{len(erros)} erro(s). Nada foi gerado; corrija o JSON e rode de novo.")
        return 1
    if args.validar:
        print(f"OK: {len(dados['casos'])} casos válidos ({len(avisos)} aviso(s)).")
        return 0

    if args.registrar:
        xlsx = Path(args.saida) / f"{slugify(dados['funcionalidade'])}_casos-de-teste.xlsx"
        if not xlsx.exists():
            print(f"ERRO: planilha não encontrada: {xlsx}. Gere os casos antes de registrar a execução.", file=sys.stderr)
            return 2
        try:
            resultados = json.loads(Path(args.registrar).read_text(encoding="utf-8"))
        except (FileNotFoundError, json.JSONDecodeError) as e:
            print(f"ERRO: não consegui ler {args.registrar}: {e}", file=sys.stderr)
            return 2
        erros_reg, contagem = registrar(dados, xlsx, resultados)
        for e in erros_reg:
            print(f"ERRO: {e}")
        if erros_reg:
            print("Nada foi gravado na planilha.")
            return 1
        resumo = ", ".join(f"{n} {s}" for s, n in contagem.items() if n)
        faltam = len(dados["casos"]) - sum(contagem.values())
        print(f"OK: execução registrada em {xlsx.name}: {resumo}" + (f"; {faltam} sem resultado." if faltam else "."))
        return 0

    saida = Path(args.saida)
    saida.mkdir(parents=True, exist_ok=True)
    base = saida / f"{slugify(dados['funcionalidade'])}_casos-de-teste"
    md, xlsx = base.with_suffix(".md"), base.with_suffix(".xlsx")
    md.write_text(gerar_md(dados), encoding="utf-8")
    gerar_xlsx(dados, xlsx)
    print(f"OK: {len(dados['casos'])} casos, {len(avisos)} aviso(s).")
    print(f"  {md.resolve()}")
    print(f"  {xlsx.resolve()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
